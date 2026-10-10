# SPDX-License-Identifier: Apache-2.0
"""Create -> patch -> re-read round trip on an in-memory DOCX. Run: python -m unittest -v"""
import io
import unittest

from docx import Document

from docx_documents import create_report, extract_text, replace_placeholders


class DocxRoundTrip(unittest.TestCase):
    def test_create_replace_extract(self) -> None:
        created = create_report(
            "Quarterly report", "Prepared for {{customer}}.", [["Item", "Qty"], ["Widgets", "3"]]
        )
        self.assertEqual(created[:2], b"PK")
        self.assertIn("{{customer}}", extract_text(created))

        patched = replace_placeholders(created, {"customer": "Ada <Lovelace> & Co"})
        text = extract_text(patched)
        self.assertIn("Quarterly report", text)
        self.assertIn("Prepared for Ada <Lovelace> & Co.", text)
        self.assertNotIn("{{customer}}", text)
        self.assertIn("Widgets\t3", text)

    def test_placeholder_split_across_runs(self) -> None:
        doc = Document()
        para = doc.add_paragraph()
        for part in ("Dear {{cus", "tomer", "}},"):
            para.add_run(part)
        buf = io.BytesIO()
        doc.save(buf)
        patched = replace_placeholders(buf.getvalue(), {"customer": "Grace"})
        self.assertIn("Dear Grace,", extract_text(patched))


if __name__ == "__main__":
    unittest.main()
