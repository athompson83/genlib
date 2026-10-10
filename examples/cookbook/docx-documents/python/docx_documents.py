# SPDX-License-Identifier: Apache-2.0
"""Genlib cookbook: create, patch and read DOCX files with python-docx (MIT).

Thin wrapper only: python-docx owns the OOXML format. Values are set as run text, so user
input is XML-escaped by the library and cannot inject markup.
"""
from __future__ import annotations

import io
from collections.abc import Iterator, Mapping, Sequence

from docx import Document
from docx.document import Document as DocumentObject
from docx.text.paragraph import Paragraph


def create_report(title: str, body: str, rows: Sequence[Sequence[str]]) -> bytes:
    """Create a DOCX with a heading, a paragraph and a table (first row = header)."""
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(body)
    table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    table.style = "Table Grid"
    for r, cells in enumerate(rows):
        for c, text in enumerate(cells):
            cell = table.cell(r, c)
            cell.text = text
            if r == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _paragraphs(doc: DocumentObject) -> Iterator[Paragraph]:
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs


def replace_placeholders(data: bytes, values: Mapping[str, str]) -> bytes:
    """Replace ``{{key}}`` placeholders in body, table cells, headers and footers.

    A placeholder inside one run keeps that run's formatting. If Word split it across
    runs, the paragraph's runs are merged into the first run (its formatting wins).
    Nested tables and text boxes are not visited.
    """
    doc = Document(io.BytesIO(data))
    tokens = {"{{" + k + "}}": v for k, v in values.items()}
    for para in _paragraphs(doc):
        if not any(t in para.text for t in tokens):
            continue
        for run in para.runs:
            for token, value in tokens.items():
                if token in run.text:
                    run.text = run.text.replace(token, value)
        if any(t in para.text for t in tokens) and para.runs:
            merged = para.text
            for token, value in tokens.items():
                merged = merged.replace(token, value)
            para.runs[0].text = merged
            for run in para.runs[1:]:
                run.text = ""
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def extract_text(data: bytes) -> str:
    """Plain text of body paragraphs and table cells, one paragraph per line."""
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(lines)
